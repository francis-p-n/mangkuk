"""Attachment text extraction and document-type identification.

Two jobs, both deliberately dumb and deterministic:

1. Turn an attachment into plain text, whatever the container format.
2. Say what the document actually *is* — because five attachments in the
   bundle are named `_SI`/`_BL` but contain a packing list, a certificate of
   origin or a commercial invoice. Trusting the filename means confidently
   extracting seven shipment fields out of a packing list.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from enum import Enum
from pathlib import Path

from .mailsource import MailSource


class DocType(str, Enum):
    SHIPPING_INSTRUCTION = "SHIPPING_INSTRUCTION"
    BILL_OF_LADING = "BILL_OF_LADING"
    PACKING_LIST = "PACKING_LIST"
    CERTIFICATE_OF_ORIGIN = "CERTIFICATE_OF_ORIGIN"
    COMMERCIAL_INVOICE = "COMMERCIAL_INVOICE"
    UNKNOWN = "UNKNOWN"


# Order matters. A PDF shipping instruction is titled "BILL OF LADING
# INSTRUCTION", which contains "bill of lading" — so the instruction markers
# must be tested before the bill-of-lading marker or every PDF SI reads as a BL.
TITLE_MARKERS: list[tuple[DocType, tuple[str, ...]]] = [
    (DocType.SHIPPING_INSTRUCTION,
     ("shipping instruction", "bill of lading instruction", "bl instruction")),
    (DocType.BILL_OF_LADING, ("bill of lading",)),
    (DocType.PACKING_LIST, ("packing list",)),
    (DocType.CERTIFICATE_OF_ORIGIN, ("certificate of origin",)),
    (DocType.COMMERCIAL_INVOICE, ("commercial invoice",)),
]

SUPPORTED = {".txt", ".xlsx", ".docx", ".pdf"}


@dataclass
class Document:
    """An attachment after extraction. `ok` false means do not trust `text`."""

    path: str
    role: str                       # "SI" or "BL" — what the email claimed it is
    fmt: str                        # file extension
    text: str = ""
    doc_type: DocType = DocType.UNKNOWN
    ok: bool = False
    # "unreadable" - the bytes arrived and could not be parsed.
    # "fetch_failed" - the bytes never arrived. Not the same thing:
    # one is a fact about the document, the other about the network,
    # and only the first is a real answer.
    error: str | None = None

    @property
    def expected_type(self) -> DocType:
        return DocType.SHIPPING_INSTRUCTION if self.role == "SI" else DocType.BILL_OF_LADING

    @property
    def type_matches_role(self) -> bool:
        return self.doc_type == self.expected_type


def identify(text: str) -> DocType:
    """Document type from its title block — the first few non-empty lines."""
    head = "\n".join(
        ln.strip() for ln in text.splitlines() if ln.strip() and set(ln.strip()) != {"="}
    )[:400].lower()
    for doc_type, markers in TITLE_MARKERS:
        if any(m in head for m in markers):
            return doc_type
    return DocType.UNKNOWN


def _from_xlsx(raw: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if not cells:
                continue
            # A two-column sheet is really a label/value pair list.
            lines.append(f"{cells[0]}: {' '.join(cells[1:])}" if len(cells) > 1 else cells[0])
    wb.close()
    return "\n".join(lines)


def _from_docx(raw: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(raw))
    lines = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if not cells:
                continue
            lines.append(f"{cells[0]}: {' '.join(cells[1:])}" if len(cells) > 1 else cells[0])
    return "\n".join(lines)


def _from_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract(source: MailSource, att_path: str) -> Document:
    """Read one attachment into a Document. Never raises — failures set `error`."""
    role = "SI" if "_SI." in att_path else "BL"
    fmt = Path(att_path).suffix.lower()
    doc = Document(path=att_path, role=role, fmt=fmt)

    if fmt not in SUPPORTED:
        # A declared-unreadable document is a correct answer; a hallucinated
        # one is not.
        doc.error = "unreadable"
        return doc

    try:
        raw = source.read_bytes(att_path)
    except Exception:
        # The document is not unreadable; we simply never saw it. Reporting
        # this as "unreadable" would turn a network blip into a confident
        # statement about a document, and quietly drop any defect in it.
        doc.error = "fetch_failed"
        return doc

    try:
        if fmt == ".txt":
            doc.text = raw.decode("utf-8", errors="replace")
        elif fmt == ".xlsx":
            doc.text = _from_xlsx(raw)
        elif fmt == ".docx":
            doc.text = _from_docx(raw)
        else:
            # Image-only or truncated PDFs extract to nothing and fall through
            # to the empty-text check below rather than being guessed at.
            doc.text = _from_pdf(raw)
    except Exception:
        doc.error = "unreadable"
        return doc

    if not doc.text.strip():
        doc.error = "unreadable"
        return doc

    doc.doc_type = identify(doc.text)
    doc.ok = True
    return doc
